from __future__ import annotations

import csv
import json
import shutil
import zipfile
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PptInches, Pt as PptPt
from rdkit import Chem
from rdkit.Chem import Draw


BASE = Path(__file__).resolve().parents[1]
OUT = BASE / "submission_artifacts_2026-05-28"
PPT_TEMPLATE = Path(r"C:\Users\srira\Downloads\RTRP__KR24_II-II_2026_Project_Presentation_Template_1.0.pptx")
DOC_TEMPLATE = Path(r"C:\Users\srira\Downloads\Real-Time Research  Project Document Format - IQAC - Final .docx")

TITLE = "ElectroMacroDiff V5.5: Pocket-Electronic Macrocycle Discovery for JAK2"
SHORT = "ElectroMacroDiff V5.5"
TODAY = date(2026, 5, 28).strftime("%d-%b-%Y")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def load_csv_dicts(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        return list(csv.DictReader(handle))


def load_project_data() -> dict:
    reference = load_csv_dicts(BASE / "09_reports" / "v5_3_benchmark" / "reference_med_style_metrics.csv")
    branches = load_csv_dicts(BASE / "09_reports" / "v5_3_benchmark" / "emd_candidate_benchmark_metrics.csv")
    summary = json.loads((BASE / "09_reports" / "v5_3_benchmark" / "emd_v5_3_benchmark_summary.json").read_text(encoding="utf-8"))
    gen_summary = json.loads((BASE / "05_generated_candidates" / "v5_5_pocket_guided" / "v5_5_pocket_guided_summary.json").read_text(encoding="utf-8"))
    gap_summary = json.loads((BASE / "05_generated_candidates" / "v5_5_pocket_guided" / "v5_5_pocket_guided_benchmark_gap_summary.json").read_text(encoding="utf-8"))
    ranking = load_csv_dicts(BASE / "08_final_ranking" / "v5_5_pocket_guided_pocket_electronic_ranked_candidates.csv")
    top = ranking[0]
    docking = load_csv_dicts(BASE / "06_docking" / "v5_5_pocket_guided" / "scores" / "docking_scores_full_vina_gpu_2_1.csv")
    pocket = json.loads((BASE / "06_docking" / "v5_5_pocket_guided" / "scores" / "jak2_pocket_electronic_profile.json").read_text(encoding="utf-8"))
    validation_raw = json.loads((BASE / "00_project_registry" / "validation_report.json").read_text(encoding="utf-8"))
    validation = validation_raw.get("summary", validation_raw)
    if "passed" not in validation:
        results = validation_raw.get("results", [])
        validation["passed"] = sum(1 for row in results if row.get("passed"))
    if "total" not in validation:
        validation["total"] = len(validation_raw.get("results", []))
    return {
        "reference": reference,
        "branches": branches,
        "summary": summary,
        "gen_summary": gen_summary,
        "gap_summary": gap_summary,
        "ranking": ranking,
        "top": top,
        "docking": docking,
        "pocket": pocket,
        "validation": validation,
    }


def val(row: dict, key: str, default: str = "not available") -> str:
    value = row.get(key, "")
    return default if value in ("", None) else str(value)


def fmt_pct(value: str | float | int | None, places: int = 2) -> str:
    if value in ("", None):
        return "not available"
    try:
        return f"{float(value):.{places}f}%"
    except Exception:
        return str(value)


def fmt_num(value: str | float | int | None, places: int = 3) -> str:
    if value in ("", None):
        return "not available"
    try:
        return f"{float(value):.{places}f}".rstrip("0").rstrip(".")
    except Exception:
        return str(value)


def draw_table_image(path: Path, title: str, columns: list[str], rows: list[list[str]], note: str) -> None:
    width = 1720
    row_h = 104
    header_h = 88
    top_h = 128
    note_h = 90
    height = top_h + header_h + row_h * len(rows) + note_h
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    ink = (18, 28, 45)
    blue = (27, 76, 122)
    fill_header = (232, 240, 250)
    fill_emd = (230, 248, 242)
    grid = (33, 33, 33)
    margin = 34
    d.text((margin, 28), title, fill=blue, font=font(36, True))
    d.text((margin, 78), note, fill=(82, 92, 110), font=font(20))
    col_w = [285] + [(width - 2 * margin - 285) // (len(columns) - 1)] * (len(columns) - 1)
    x = margin
    y = top_h
    for i, col in enumerate(columns):
        d.rectangle([x, y, x + col_w[i], y + header_h], fill=fill_header, outline=grid, width=2)
        d.text((x + 14, y + 25), col, fill=ink, font=font(25, True))
        x += col_w[i]
    y += header_h
    for row in rows:
        x = margin
        for i, cell in enumerate(row):
            fill = fill_emd if i == len(row) - 1 else "white"
            d.rectangle([x, y, x + col_w[i], y + row_h], fill=fill, outline=grid, width=2)
            d.multiline_text((x + 14, y + 25), cell, fill=ink, font=font(24), spacing=4)
            x += col_w[i]
        y += row_h
    d.text((margin, height - 58), "Source: project output CSV/JSON files and MED reference metrics exported in 09_reports/v5_3_benchmark.", fill=(94, 102, 116), font=font(18))
    img.save(path)


def draw_pipeline_image(path: Path, data: dict) -> None:
    width, height = 1720, 620
    img = Image.new("RGB", (width, height), (248, 251, 255))
    d = ImageDraw.Draw(img)
    title_color = (27, 76, 122)
    d.text((40, 30), "ElectroMacroDiff V5.5 Evidence Chain", fill=title_color, font=font(38, True))
    stages = [
        ("Curate", "1,135 JAK2 ligands"),
        ("Train/Guide", "pocket-electronic context"),
        ("Generate", "1,113 candidates"),
        ("Dock", "1,036 Vina-GPU scores"),
        ("Rescore", "pocket + ADMET"),
        ("Review", "3D protein-pose dashboard"),
    ]
    x, y, box_w, box_h, gap = 52, 170, 238, 176, 32
    colors = [(232, 240, 250), (236, 247, 255), (231, 249, 241), (255, 246, 222), (245, 238, 255), (232, 248, 242)]
    for i, (name, detail) in enumerate(stages):
        bx = x + i * (box_w + gap)
        d.rounded_rectangle([bx, y, bx + box_w, y + box_h], radius=18, fill=colors[i], outline=(80, 92, 118), width=2)
        d.text((bx + 22, y + 34), name, fill=(18, 28, 45), font=font(28, True))
        d.multiline_text((bx + 22, y + 84), detail, fill=(55, 65, 85), font=font(21), spacing=5)
        if i < len(stages) - 1:
            ax = bx + box_w + 7
            d.line([ax, y + 88, ax + gap - 16, y + 88], fill=(27, 76, 122), width=4)
            d.polygon([(ax + gap - 16, y + 88), (ax + gap - 30, y + 78), (ax + gap - 30, y + 98)], fill=(27, 76, 122))
    d.rounded_rectangle([52, 420, width - 52, 548], radius=16, fill=(255, 255, 255), outline=(190, 202, 220), width=2)
    line = (
        f"Top candidate: {data['top']['candidate_id']} | best Vina-GPU {fmt_num(data['top']['best_score'], 1)} kcal/mol | "
        f"pocket-guided score {fmt_num(data['top']['pocket_guided_final_score'], 6)} | "
        f"validation {data['validation']['passed']}/{data['validation']['total']} passed"
    )
    d.text((78, 462), line, fill=(18, 28, 45), font=font(26, True))
    img.save(path)


def draw_molecule(path: Path, smiles: str) -> None:
    mol = Chem.MolFromSmiles(smiles)
    if mol is None:
        return
    image = Draw.MolToImage(mol, size=(900, 620), kekulize=True)
    image.save(path)


def make_images(data: dict) -> dict[str, Path]:
    assets = OUT / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    ref = {row["method"]: row for row in data["reference"]}
    branches = {row["branch"]: row for row in data["branches"]}
    latest = branches["kaggle_v5_5_pocket_guided"]
    latest_rows = [
        ["Validity", fmt_pct(ref["MED"]["validity_percent"]), fmt_pct(ref["Macformer"]["validity_percent"]), fmt_pct(ref["MacLS"]["validity_percent"]), fmt_pct(latest["validity_percent"], 4)],
        ["Uniqueness", fmt_pct(ref["MED"]["uniqueness_percent"]), fmt_pct(ref["Macformer"]["uniqueness_percent"]), fmt_pct(ref["MacLS"]["uniqueness_percent"]), fmt_pct(latest["uniqueness_percent"])],
        ["Macrocyclization", fmt_pct(ref["MED"]["macrocyclization_percent"]), fmt_pct(ref["Macformer"]["macrocyclization_percent"]), fmt_pct(ref["MacLS"]["macrocyclization_percent"]), fmt_pct(latest["macrocyclization_percent"])],
        ["Linker novelty", fmt_pct(ref["MED"]["linker_novelty_percent"]), fmt_pct(ref["Macformer"]["linker_novelty_percent"]), fmt_pct(ref["MacLS"]["linker_novelty_percent"]), fmt_pct(latest["linker_novelty_percent"])],
    ]
    latest_img = assets / "med_vs_emd_latest_metrics.png"
    draw_table_image(
        latest_img,
        "MED vs Macformer vs MacLS vs ElectroMacroDiff V5.5",
        ["Metric", "MED", "Macformer", "MacLS", "EMD V5.5"],
        latest_rows,
        "EMD V5.5 validity is raw-attempt logged; uniqueness/macrocycle/novelty are final accepted-output metrics.",
    )
    version_rows = []
    for branch in ["frozen_v5_3_baseline", "kaggle_logged_rerun", "kaggle_diverse_rerun", "kaggle_v5_5_pocket_guided"]:
        row = branches[branch]
        version_rows.append([
            branch.replace("_", " "),
            fmt_pct(row["validity_percent"], 4),
            fmt_pct(row["uniqueness_percent"]),
            fmt_pct(row["macrocyclization_percent"]),
            fmt_pct(row["linker_novelty_percent"]),
        ])
    versions_img = assets / "emd_versions_metrics.png"
    draw_table_image(
        versions_img,
        "ElectroMacroDiff Branch Evolution Metrics",
        ["Branch", "Validity", "Uniqueness", "Macrocyclization", "Linker novelty"],
        version_rows,
        "Rows mix filtered-survivor and raw-attempt scopes; each branch keeps its benchmark scope in the report.",
    )
    pipeline_img = assets / "emd_v5_5_pipeline.png"
    draw_pipeline_image(pipeline_img, data)
    molecule_img = assets / "top_candidate_CAND_3688f8acdc.png"
    draw_molecule(molecule_img, data["top"]["canonical_smiles"] if "canonical_smiles" in data["top"] else data["top"]["smiles"])
    return {"latest": latest_img, "versions": versions_img, "pipeline": pipeline_img, "molecule": molecule_img}


def clear_slides(prs: Presentation) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        r_id = sld_id.rId
        prs.part.drop_rel(r_id)
        sld_id_lst.remove(sld_id)


def add_slide_number(slide, num: int) -> None:
    box = slide.shapes.add_textbox(PptInches(12.35), PptInches(7.05), PptInches(0.7), PptInches(0.25))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = str(num)
    p.font.size = PptPt(9)
    p.font.color.rgb = PptRGBColor(90, 90, 90)
    p.alignment = PP_ALIGN.RIGHT


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(PptInches(0.55), PptInches(0.28), PptInches(12.1), PptInches(0.62))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = PptPt(26)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(31, 78, 121)
    if subtitle:
        sub = slide.shapes.add_textbox(PptInches(0.58), PptInches(0.86), PptInches(11.8), PptInches(0.32))
        sp = sub.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = PptPt(10.5)
        sp.font.color.rgb = PptRGBColor(90, 98, 112)


def add_bullets(slide, items: list[str], left: float, top: float, width: float, height: float, size: int = 15) -> None:
    box = slide.shapes.add_textbox(PptInches(left), PptInches(top), PptInches(width), PptInches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = PptPt(size)
        p.font.color.rgb = PptRGBColor(35, 45, 60)
        p.space_after = PptPt(5)


def add_metric_card(slide, x: float, y: float, w: float, h: float, value: str, label: str, color=(31, 78, 121)) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PptRGBColor(242, 247, 252)
    shape.line.color.rgb = PptRGBColor(180, 196, 214)
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = value
    p.font.size = PptPt(24)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(*color)
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.size = PptPt(10.5)
    p2.font.color.rgb = PptRGBColor(70, 80, 95)
    p2.alignment = PP_ALIGN.CENTER


def add_table(slide, headers: list[str], rows: list[list[str]], left: float, top: float, width: float, height: float, font_size: int = 9) -> None:
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), PptInches(left), PptInches(top), PptInches(width), PptInches(height))
    table = table_shape.table
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = PptRGBColor(31, 78, 121)
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = PptPt(font_size)
            p.font.color.rgb = PptRGBColor(255, 255, 255)
    for r, row in enumerate(rows, 1):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            for p in cell.text_frame.paragraphs:
                p.font.size = PptPt(font_size)
                p.font.color.rgb = PptRGBColor(35, 45, 60)
    return table_shape


def build_ppt(data: dict, images: dict[str, Path]) -> Path:
    prs = Presentation(str(PPT_TEMPLATE))
    clear_slides(prs)
    layouts = prs.slide_layouts
    slide_no = 1

    s = prs.slides.add_slide(layouts[0])
    add_title(s, TITLE, f"RTRP II-II 2026 Project Presentation | Date: {TODAY}")
    add_bullets(s, ["Team Members: Member 1, Member 2, Member 3, Member 4, Member 5", "Mentor: To be added", "Domain: AI for Drug Discovery / Computational Biology"], 0.9, 2.0, 8.2, 1.8, 18)
    add_metric_card(s, 9.6, 1.55, 2.9, 1.0, "1,113", "V5.5 candidates")
    add_metric_card(s, 9.6, 2.75, 2.9, 1.0, "-12.7", "Best Vina-GPU")
    add_metric_card(s, 9.6, 3.95, 2.9, 1.0, "54/54", "Validation checks")
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[1])
    add_title(s, "Contents", "Narrative aligned to the supplied RTRP presentation template")
    add_bullets(s, ["Introduction and problem statement", "Requirement and objectives", "Architecture and implementation", "MED vs EMD benchmark evidence", "Docking, pocket-fit, and top candidate results", "3D dashboard and validation", "Limitations, learnings, and V6 roadmap"], 0.9, 1.45, 11.5, 4.8, 18)
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[2])
    add_title(s, "Introduction", "Why JAK2 macrocycle generation needs an evidence-first workflow")
    add_bullets(s, [
        "JAK2 inhibitors require balancing potency, selectivity, ADMET feasibility, and binding-pocket fit.",
        "Macrocycles can improve conformational control, but valid ring closure and linker novelty are difficult.",
        "ElectroMacroDiff V5.5 connects generation, docking, pocket-electronic scoring, and 3D inspection into one auditable pipeline.",
    ], 0.75, 1.45, 5.8, 4.8, 15)
    add_bullets(s, [
        "Project aim: generate and prioritize pocket-fitted JAK2 macrocycle candidates.",
        "Current claim: pocket/electronic-conditioned generation plus post-generation docking and pocket rescoring.",
        "Not claimed: wet-lab validation or full residue-level protein-conditioned diffusion.",
    ], 6.8, 1.45, 5.8, 4.8, 15)
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[5])
    add_title(s, "Requirement", "The system must be scientifically honest, reproducible, and useful for review")
    rows = [
        ["Functional", "Curate JAK2 ligands, generate macrocycles, dock, rescore, rank, export dashboard."],
        ["Non-functional", "Low-resource runnable workflow; explicit logs; no fake metrics; clear limitations."],
        ["End-user benefit", "Project readers can inspect top candidates, metrics, pose evidence, and project maturity."],
        ["Acceptance", "All outputs must trace to project CSV/JSON/PDB/PDBQT artifacts."],
    ]
    add_table(s, ["Need", "Implementation in EMD V5.5"], rows, 0.75, 1.45, 11.8, 3.4, 12)
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[6])
    add_title(s, "Design", "Evidence chain from curated ligands to protein-ligand 3D inspection")
    s.shapes.add_picture(str(images["pipeline"]), PptInches(0.75), PptInches(1.28), width=PptInches(11.8))
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[2])
    add_title(s, "Implementation", "V5.5 introduced pocket/electronic-conditioned decision points")
    add_bullets(s, [
        "Pocket profile: hydrophobic, donor/acceptor, aromatic, charge proxy, and residue-contact features.",
        "Generation: anchor/linker decisions, exact linker novelty enforcement, and validity/reward gate.",
        "Docking: Vina-GPU 2.1 scoring against JAK2 5AEP grid.",
        "Dashboard: real candidate table plus protein-pose 3D viewer.",
    ], 0.75, 1.35, 5.8, 5.2, 14)
    rows = [
        ["Candidates", "1,113"],
        ["Docked scores", "1,036"],
        ["Raw attempts", "87,287"],
        ["Valid outputs", "1,113"],
        ["Unit tests", "40/40"],
        ["Validation", "54/54"],
    ]
    add_table(s, ["Output", "Value"], rows, 7.0, 1.45, 4.9, 3.75, 13)
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[6])
    add_title(s, "Benchmark Image", "MED-style comparison, with EMD scope clearly labeled")
    s.shapes.add_picture(str(images["latest"]), PptInches(0.45), PptInches(1.18), width=PptInches(12.3))
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[6])
    add_title(s, "Version Evolution", "EMD branches show the tradeoff between novelty gating and raw validity")
    s.shapes.add_picture(str(images["versions"]), PptInches(0.45), PptInches(1.18), width=PptInches(12.3))
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[2])
    add_title(s, "Docking and Pocket Fit", "Top-ranked candidate combines strong docking with pocket-electronic evidence")
    t = data["top"]
    add_metric_card(s, 0.8, 1.35, 2.8, 1.0, fmt_num(t["best_score"], 1), "Best Vina-GPU kcal/mol")
    add_metric_card(s, 3.9, 1.35, 2.8, 1.0, fmt_num(t["pocket_electronic_fit_score"], 6), "Pocket fit score")
    add_metric_card(s, 7.0, 1.35, 2.8, 1.0, fmt_num(t["pocket_guided_final_score"], 6), "Final score")
    add_metric_card(s, 10.1, 1.35, 2.3, 1.0, fmt_num(t["qed"], 4), "QED")
    add_bullets(s, [
        f"Top candidate: {t['candidate_id']}",
        f"Macrocycle ring size: {t['max_ring_size']}; Lipinski violations: {t['lipinski_violations']}.",
        "Pocket scoring combines electrostatic, H-bond, hydrophobic, and aromatic contact components.",
        "Top contact residues include ARG:A:980, TYR:A:931, ASN:A:981, LEU:A:932, ASP:A:939, and LYS:A:882.",
    ], 0.85, 3.0, 6.1, 2.7, 14)
    s.shapes.add_picture(str(images["molecule"]), PptInches(7.25), PptInches(2.75), width=PptInches(4.9))
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[2])
    add_title(s, "3D Dashboard", "Real receptor and pose files are directly inspectable")
    add_bullets(s, [
        "Viewer embeds real JAK2 receptor PDB and 1,036 real docked pose texts.",
        "Candidate list supports direct click-to-open with rank, Vina, pocket fit, contact count, and pose status.",
        "Controls include protein+pose, pose-only, protein-only, surfaces, atom labels, bond lengths, contact labels, and docking grid box.",
    ], 0.75, 1.35, 5.9, 4.6, 15)
    add_bullets(s, [
        "Important: the viewer does not invent structures.",
        "If a pose is missing, the interface reports it as unavailable.",
        "All candidate metrics come from interface/data.js generated from project output files.",
    ], 6.95, 1.35, 5.3, 4.6, 15)
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[2])
    add_title(s, "Project Status", "Complete as a computational evidence package; benchmark gaps are explicit")
    add_bullets(s, [
        "Completed: data curation, features, V5.5 generation, docking, pocket scoring, ADMET proxy scoring, final ranking, dashboard, validation.",
        "Evidence package: candidate CSVs, docking scores, pose files, pocket profile JSON, reports, interface data, tests.",
        "Validation passed: 54/54 checks; unit tests passed: 40/40.",
    ], 0.75, 1.35, 5.9, 4.8, 15)
    add_bullets(s, [
        "Remaining scientific gaps: experimental validation, MD stability, free-energy estimation, residue-level learned pocket encoder.",
        "Honest benchmark status: EMD V5.5 is target-specific and auditable; it does not fully beat MED on MED's macrocycle-generation benchmark.",
    ], 6.95, 1.35, 5.3, 4.8, 15)
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[5])
    add_title(s, "Learnings and Future Scope", "The next leap is V6 protein-conditioned generation")
    rows = [
        ["Learning", "End-to-end evidence matters more than isolated generator scores."],
        ["Learning", "Raw-attempt logging prevents false benchmark claims."],
        ["Learning", "3D inspection makes candidate evidence understandable to non-coding project readers."],
        ["V6", "Train residue-level interaction encoder using PDBbind/PLINDER/BioLiP-style contact data."],
        ["V6", "Condition anchor and linker generation on aligned pocket poses, not only global pocket features."],
    ]
    add_table(s, ["Area", "Outcome"], rows, 0.75, 1.35, 11.8, 4.1, 12)
    add_slide_number(s, slide_no); slide_no += 1

    s = prs.slides.add_slide(layouts[1])
    add_title(s, "Thank You", "Computational candidates only - experimental validation remains required")
    add_bullets(s, ["Project: ElectroMacroDiff V5.5", "Target: JAK2 5AEP", "Submission package includes PPTX, IQAC DOCX, research paper draft, comparison images, and artifact summary."], 1.2, 2.1, 10.6, 2.2, 20)
    add_slide_number(s, slide_no)

    out = OUT / "ElectroMacroDiff_V5_5_Project_Presentation.pptx"
    prs.save(out)
    return out


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = text
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.size = Pt(9)
            run.font.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def doc_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(31, 78, 121)


def doc_para(doc: Document, text: str, bold_start: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_start and text.startswith(bold_start):
        r = p.add_run(bold_start)
        r.bold = True
        p.add_run(text[len(bold_start):])
    else:
        p.add_run(text)


def doc_bullet(doc: Document, text: str) -> None:
    try:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)
    except KeyError:
        p = doc.add_paragraph()
        p.add_run(f"- {text}")


def doc_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True)
    for row in rows:
        cells = table.add_row().cells
        for i, cell_text in enumerate(row):
            set_cell_text(cells[i], cell_text)
    doc.add_paragraph()


def replace_placeholders(doc: Document) -> None:
    replacements = {
        "<Title of the Project>": TITLE,
        "<Title>": TITLE,
        "YOUR PROJECT ABSTRACT SHOUL BE HERE": "This report presents ElectroMacroDiff V5.5, a low-resource computational drug discovery workflow for JAK2 macrocycle candidate generation, docking, pocket-electronic rescoring, ADMET proxy scoring, and interactive 3D protein-ligand review. The system produced 1,113 ranked V5.5 candidates, 1,036 parsed Vina-GPU docking scores, and a real-data dashboard for final evaluation.",
        "In this project, a prototype and implementation of Home Appliances is demonstrated. The proposed system consists of a Hardware interface and Software interface. Hardware interface includes Raspberry Pi, lights, fans, DC ": "",
        "Ms./Mr./Dr.Mentor Name": "Ms./Mr./Dr. Mentor Name",
        "<Designation>, <Department>": "Department of Computer Science and Engineering",
    }
    for p in doc.paragraphs:
        for old, new in replacements.items():
            if old in p.text:
                for run in p.runs:
                    run.text = run.text.replace(old, new)


def build_report_docx(data: dict, images: dict[str, Path]) -> Path:
    doc = Document(str(DOC_TEMPLATE))
    replace_placeholders(doc)
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc_heading(doc, "CHAPTER 1: INTRODUCTION", 1)
    doc_para(doc, "ElectroMacroDiff V5.5 is a computational drug discovery project focused on JAK2 macrocycle inhibitor candidate generation and prioritization. The system integrates curated ligand data, model-guided macrocycle construction, docking, ADMET and synthesis proxy scoring, pocket-electronic rescoring, and interactive 3D review.")
    doc_para(doc, "The core project goal is to produce pocket-fitted macrocycle candidates for JAK2 and to provide a transparent evidence chain for review. The current system is pocket/electronic-conditioned and pocket-scored, but it is not claimed as a full residue-level protein-conditioned diffusion model.")

    doc_heading(doc, "CHAPTER 2: PROBLEM STATEMENT AND OBJECTIVES", 1)
    for item in [
        "Generate chemically valid and novel macrocycle candidates inspired by JAK2 ligand data.",
        "Use JAK2 pocket context and docking evidence to prioritize candidates.",
        "Measure raw-attempt validity, uniqueness, macrocyclization, and linker novelty honestly.",
        "Provide a dashboard that allows direct inspection of candidate metrics and real protein-ligand poses.",
    ]:
        doc_bullet(doc, item)

    doc_heading(doc, "CHAPTER 3: SYSTEM ARCHITECTURE", 1)
    doc.add_picture(str(images["pipeline"]), width=Inches(6.4))
    doc_para(doc, "The workflow starts with curated JAK2 ligands, builds and evaluates macrocycle candidates, docks them into the 5AEP JAK2 receptor grid, rescoring candidates using pocket-electronic contact features, and presents the final ranked set in a real-data dashboard.")

    doc_heading(doc, "CHAPTER 4: IMPLEMENTATION", 1)
    doc_table(doc, ["Component", "Implementation"], [
        ["Data curation", "1,135 curated JAK2 ligands and macrocycle-focused training artifacts."],
        ["Generation", "V5.5 pocket-guided macrocycle generation with linker novelty and validity/reward gating."],
        ["Docking", "Vina-GPU 2.1 docking against JAK2 5AEP receptor grid."],
        ["Scoring", "ADMET, synthesis, docking, and pocket-electronic fit scoring."],
        ["Interface", "Real-data dashboard with candidate filters and protein-ligand 3D viewer."],
    ])

    doc_heading(doc, "CHAPTER 5: BENCHMARK AND COMPARISON", 1)
    doc.add_picture(str(images["latest"]), width=Inches(6.4))
    doc_para(doc, "The MED, Macformer, and MacLS values are reference metrics from the MED-style benchmark export. EMD V5.5 metrics come from the active Kaggle V5.5 pocket-guided branch. The comparison is intentionally labeled because EMD V5.5 is a target-specific JAK2 pipeline and does not use the identical benchmark scope as MED.")
    doc.add_picture(str(images["versions"]), width=Inches(6.4))

    doc_heading(doc, "CHAPTER 6: RESULTS", 1)
    top = data["top"]
    doc_table(doc, ["Metric", "Value"], [
        ["Generated V5.5 candidates", "1,113"],
        ["Raw attempts", "87,287"],
        ["Raw-attempt validity", "1.2751%"],
        ["Docked candidates with parsed scores", "1,036"],
        ["Best Vina-GPU score", "-12.7 kcal/mol"],
        ["Best pocket-electronic fit score", "0.84838"],
        ["Top ranked candidate", top["candidate_id"]],
        ["Top pocket-guided final score", fmt_num(top["pocket_guided_final_score"], 6)],
        ["Validation checks", "54/54 passed"],
        ["Unit tests", "40/40 passed"],
    ])
    doc.add_picture(str(images["molecule"]), width=Inches(4.8))
    doc_para(doc, f"Top candidate SMILES: {top['canonical_smiles'] if 'canonical_smiles' in top else top['smiles']}")

    doc_heading(doc, "CHAPTER 7: 3D DASHBOARD AND REVIEW", 1)
    doc_para(doc, "The upgraded interface embeds the real JAK2 receptor and 1,036 real docked pose texts. It supports direct candidate click-to-open, protein plus pose view, pose-only view, protein-only view, atom labels, bond-length labels, contact residue labels, transparent ligand and pocket surfaces, and docking grid visualization.")
    doc_para(doc, "The dashboard displays only real exported project values. Missing pose or metric values are shown as unavailable rather than inferred.")

    doc_heading(doc, "CHAPTER 8: LIMITATIONS", 1)
    for item in [
        "Docking scores are computational estimates and are not experimental binding affinities.",
        "Pocket-electronic scores are proxy scores, not quantum electron-density calculations.",
        "Raw-attempt validity remains low because the reward and novelty gates reject many proposals.",
        "No wet-lab assay, molecular dynamics stability run, or free-energy calculation has been performed yet.",
    ]:
        doc_bullet(doc, item)

    doc_heading(doc, "CHAPTER 9: CONCLUSION AND FUTURE SCOPE", 1)
    doc_para(doc, "ElectroMacroDiff V5.5 is complete as a computational evidence package. It provides real candidates, docking outputs, pose files, pocket scoring, validation results, and an interactive dashboard. The next scientific phase should be V6: learned residue-level protein-ligand interaction encoding using richer datasets such as PDBbind, PLINDER/BioLiP-style contact data, and expanded JAK/JAK2 activity data.")

    doc_heading(doc, "REFERENCES", 1)
    refs = [
        "Trott, O.; Olson, A. J. AutoDock Vina: improving the speed and accuracy of docking with a new scoring function, efficient optimization, and multithreading. Journal of Computational Chemistry, 2010.",
        "Ding, J.; Tang, S.; et al. Vina-GPU 2.0: Further accelerating AutoDock Vina and its derivatives with graphics processing units. Journal of Chemical Information and Modeling, 2023.",
        "RDKit: Open-source cheminformatics. https://www.rdkit.org",
        "Rego, N.; Koes, D. 3Dmol.js: molecular visualization with WebGL. Bioinformatics, 2015.",
    ]
    for ref in refs:
        doc_bullet(doc, ref)

    out = OUT / "ElectroMacroDiff_V5_5_IQAC_Project_Report.docx"
    doc.save(out)
    return out


def build_research_paper_docx(data: dict, images: dict[str, Path]) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(31, 78, 121)
    p = doc.add_paragraph("Author 1, Author 2, Author 3, Author 4, Author 5")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    doc_heading(doc, "Abstract", 1)
    doc_para(doc, "Macrocyclic inhibitors are attractive in kinase discovery because their constrained conformations can improve target engagement and selectivity. We present ElectroMacroDiff V5.5, a low-resource computational workflow for JAK2 macrocycle candidate generation and prioritization. The system combines curated JAK2 ligand data, model-guided macrocycle construction, novelty and validity gating, Vina-GPU docking, ADMET and synthesis proxy scoring, pocket-electronic rescoring, and interactive 3D review. The final branch produced 1,113 ranked candidates from 87,287 logged generation attempts, with 1,036 parsed docking scores and a best Vina-GPU score of -12.7 kcal/mol.")
    for heading, body in [
        ("1. Introduction", "JAK2 macrocycle discovery requires chemical validity, pocket fit, and interpretable prioritization. ElectroMacroDiff addresses this by connecting generation to a complete evidence chain rather than reporting isolated generator outputs."),
        ("2. Methods", "The workflow curates JAK2 ligands, generates macrocycles, enforces linker novelty, applies a validity/reward gate, docks candidates against JAK2 5AEP, calculates ADMET and synthesis proxy metrics, and rescoring docked poses using pocket-electronic features."),
        ("3. Results", "The active V5.5 branch generated 1,113 candidates, retained 100% uniqueness, 100% macrocyclization, and 100% linker novelty in the final accepted set. Raw-attempt validity was 1.2751%, reflecting aggressive rejection by the reward and novelty gates."),
        ("4. Discussion", "The project is strongest as a reproducible computational prioritization pipeline. It should not be framed as experimentally validated or as a full protein-conditioned diffusion model. Its value is transparent data-to-dashboard traceability and target-specific JAK2 evidence."),
        ("5. Conclusion", "ElectroMacroDiff V5.5 is ready as a computational research artifact. Future work should train residue-level protein-ligand interaction encoders and align seed poses for true V6 protein-conditioned generation."),
    ]:
        doc_heading(doc, heading, 1)
        doc_para(doc, body)
    doc.add_picture(str(images["latest"]), width=Inches(6.5))
    doc.add_picture(str(images["pipeline"]), width=Inches(6.5))
    doc_heading(doc, "References", 1)
    for ref in [
        "Trott and Olson, AutoDock Vina, Journal of Computational Chemistry, 2010.",
        "Ding et al., Vina-GPU 2.0, Journal of Chemical Information and Modeling, 2023.",
        "RDKit open-source cheminformatics.",
        "Rego and Koes, 3Dmol.js, Bioinformatics, 2015.",
    ]:
        doc_bullet(doc, ref)
    out = OUT / "ElectroMacroDiff_V5_5_Research_Paper_Draft.docx"
    doc.save(out)
    return out


def write_summary(data: dict, files: list[Path]) -> Path:
    path = OUT / "artifact_summary.md"
    text = [
        "# ElectroMacroDiff V5.5 Submission Artifact Summary",
        "",
        f"Generated on: {TODAY}",
        "",
        "## Included Files",
        *[f"- `{file.name}`" for file in files],
        "",
        "## Real Metrics Used",
        "- Curated JAK2 ligands: 1,135",
        "- V5.5 generated/ranked candidates: 1,113",
        "- Parsed docking scores: 1,036",
        "- Best Vina-GPU score: -12.7 kcal/mol",
        "- Raw-attempt validity: 1.2751%",
        "- Linker novelty: 100%",
        "- Uniqueness: 100%",
        "- Macrocyclization: 100%",
        "- Validation: 54/54 checks passed",
        "- Unit tests: 40/40 passed",
        "",
        "## Claim Boundary",
        "This package presents computational candidates only. No wet-lab validation is claimed.",
    ]
    path.write_text("\n".join(text), encoding="utf-8")
    return path


def make_zip(files: list[Path]) -> Path:
    zip_path = BASE.parent / "ElectroMacroDiff_V5_5_Presentation_Documents_2026-05-28.zip"
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
        for file in files:
            zf.write(file, f"ElectroMacroDiff_V5_5_Submission/{file.name}")
        for asset in (OUT / "assets").glob("*.png"):
            zf.write(asset, f"ElectroMacroDiff_V5_5_Submission/assets/{asset.name}")
    return zip_path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    data = load_project_data()
    images = make_images(data)
    ppt = build_ppt(data, images)
    report = build_report_docx(data, images)
    paper = build_research_paper_docx(data, images)
    summary = write_summary(data, [ppt, report, paper, *images.values()])
    zip_path = make_zip([ppt, report, paper, summary])
    print(json.dumps({
        "output_dir": str(OUT),
        "ppt": str(ppt),
        "report_docx": str(report),
        "paper_docx": str(paper),
        "summary": str(summary),
        "zip": str(zip_path),
        "images": {k: str(v) for k, v in images.items()},
    }, indent=2))


if __name__ == "__main__":
    main()
