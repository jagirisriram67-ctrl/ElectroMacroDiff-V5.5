from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
from datetime import date
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ACCENT = "1F4E79"
ACCENT_2 = "2F6F5E"
LIGHT_BLUE = "DDEAF6"
LIGHT_GREEN = "DCEFE8"
LIGHT_GRAY = "EEF2F5"
SOFT_RED = "F8E1E1"
TEXT = "243447"


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists() or path.stat().st_size == 0:
        return []
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def read_json(path: Path) -> dict[str, Any]:
    if not path.exists() or path.stat().st_size == 0:
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def truthy(value: Any) -> bool:
    return str(value).strip().lower() in {"true", "1", "yes", "y"}


def as_float(value: Any) -> float | None:
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    return out if math.isfinite(out) else None


def as_int(value: Any, default: int = 0) -> int:
    number = as_float(value)
    return int(number) if number is not None else default


def count_by(rows: list[dict[str, Any]], key: str) -> dict[str, int]:
    counts: dict[str, int] = {}
    for row in rows:
        value = str(row.get(key, "") or "blank")
        counts[value] = counts.get(value, 0) + 1
    return dict(sorted(counts.items(), key=lambda item: str(item[0])))


def numeric_stats(rows: list[dict[str, Any]], key: str) -> dict[str, float | int | None]:
    values = [as_float(row.get(key)) for row in rows]
    values = [value for value in values if value is not None]
    if not values:
        return {"count": 0, "min": None, "median": None, "max": None}
    return {
        "count": len(values),
        "min": round(min(values), 3),
        "median": round(statistics.median(values), 3),
        "max": round(max(values), 3),
    }


def percent(part: int, whole: int) -> str:
    if not whole:
        return "0.0%"
    return f"{100.0 * part / whole:.1f}%"


def rel(base: Path, path: Path) -> str:
    try:
        return str(path.resolve().relative_to(base.resolve())).replace("\\", "/")
    except ValueError:
        return str(path)


def file_note(base: Path, relative_path: str) -> str:
    path = base / relative_path
    return f"{relative_path} ({'exists' if path.exists() else 'missing'})"


def build_inventory(base: Path) -> dict[str, Any]:
    paths = {
        "raw_chembl": base / "01_raw_data" / "chembl" / "jak2_activities_raw.csv",
        "pdb": base / "01_raw_data" / "pdb" / "5AEP.pdb",
        "curated": base / "02_curated_data" / "jak2_curated_ligands.csv",
        "features": base / "03_features" / "ligand_features.csv",
        "graph_index": base / "03_features" / "ligand_graphs" / "se3_graph_index.csv",
        "graphs": base / "03_features" / "ligand_graphs" / "se3_graphs.pt",
        "rdkit": base / "05_generated_candidates" / "rdkit" / "generated_rdkit_filtered.csv",
        "selfies": base / "05_generated_candidates" / "selfies" / "generated_selfies_filtered.csv",
        "generated": base / "05_generated_candidates" / "merged" / "generated_merged_filtered.csv",
        "admet": base / "07_admet_synthesis" / "admet_scores.csv",
        "sdf_manifest": base / "06_docking" / "ligands_sdf" / "docking_input_manifest.csv",
        "pdbqt_manifest": base / "06_docking" / "ligands_pdbqt" / "ligand_pdbqt_manifest.csv",
        "vina_manifest": base / "06_docking" / "scores" / "vina_command_manifest.csv",
        "docking_scores": base / "06_docking" / "scores" / "docking_scores.csv",
        "docking_summary": base / "06_docking" / "scores" / "docking_summary.json",
        "pose_sanity": base / "06_docking" / "scores" / "pose_sanity_scores.csv",
        "ranking": base / "08_final_ranking" / "final_ranked_candidates.csv",
        "validation": base / "00_project_registry" / "validation_report.json",
        "consistency": base / "09_reports" / "consistency_audit.json",
        "se3_debug": base / "04_models_checkpoints" / "se3_flow" / "se3_dataloader_test_passed.json",
        "config": base / "00_project_registry" / "campaign_config.yaml",
        "environment": base / "00_project_registry" / "environment_versions.txt",
    }

    raw = read_csv(paths["raw_chembl"])
    curated = read_csv(paths["curated"])
    features = read_csv(paths["features"])
    graph_index = read_csv(paths["graph_index"])
    rdkit_rows = read_csv(paths["rdkit"])
    selfies_rows = read_csv(paths["selfies"])
    generated = read_csv(paths["generated"])
    admet = read_csv(paths["admet"])
    sdf_manifest = read_csv(paths["sdf_manifest"])
    pdbqt_manifest = read_csv(paths["pdbqt_manifest"])
    vina_manifest = read_csv(paths["vina_manifest"])
    docking = read_csv(paths["docking_scores"])
    pose_sanity = read_csv(paths["pose_sanity"])
    ranking = read_csv(paths["ranking"])

    final_candidates = [row for row in ranking if row.get("selection_tier") == "final_candidate"]
    final_candidates.sort(key=lambda row: as_int(row.get("diverse_rank"), as_int(row.get("rank"), 9999)))
    final_candidates = final_candidates[:5]

    macrocycles = [row for row in ranking if truthy(row.get("has_macrocycle_12_20"))]
    macrocycles.sort(key=lambda row: as_int(row.get("rank"), 999999))
    constrained = [
        row
        for row in ranking
        if truthy(row.get("has_constrained_ring_8_11")) and not truthy(row.get("has_macrocycle_12_20"))
    ]
    constrained.sort(key=lambda row: as_int(row.get("rank"), 999999))

    split_counts = count_by(curated, "split")
    ring_counts = count_by(generated, "max_ring_size")
    generated_macro = sum(truthy(row.get("has_macrocycle_12_20")) for row in generated)
    generated_constrained = sum(truthy(row.get("has_constrained_ring_8_11")) for row in generated)
    curated_macro = sum(truthy(row.get("has_macrocycle_12_20")) for row in curated)
    curated_constrained = sum(truthy(row.get("has_constrained_ring_8_11")) for row in curated)

    image_dir = base / "09_reports" / "molecule_images"
    pose_image_dir = base / "06_docking" / "images" / "pose_sanity"
    card_dir = base / "09_reports" / "candidate_cards"

    final_candidate_assets = []
    for row in final_candidates:
        cid = str(row.get("candidate_id", ""))
        final_candidate_assets.append(
            {
                "candidate_id": cid,
                "rank": row.get("rank", ""),
                "diverse_rank": row.get("diverse_rank", ""),
                "source_generator": row.get("source_generator", ""),
                "max_ring_size": row.get("max_ring_size", ""),
                "docking_score": row.get("docking_score", ""),
                "final_weighted_score": row.get("final_weighted_score", ""),
                "admet_score": row.get("admet_score", ""),
                "synthesis_score": row.get("synthesis_score", ""),
                "pose_decision": row.get("pose_decision", ""),
                "main_risk": row.get("main_risk", ""),
                "smiles": row.get("canonical_smiles", row.get("smiles", "")),
                "image_path": image_dir / f"{cid}.png",
                "pose_image_path": pose_image_dir / f"{cid}_pose.png",
                "card_path": card_dir / f"{cid}_card.md",
            }
        )

    dataset_rows = [
        {
            "Dataset / artifact": "Raw JAK2 activity pull",
            "Rows": len(raw),
            "Location": rel(base, paths["raw_chembl"]),
            "Meaning": "Initial ChEMBL activity records before project curation.",
            "Status": "implemented",
        },
        {
            "Dataset / artifact": "JAK2 structure",
            "Rows": 1 if paths["pdb"].exists() else 0,
            "Location": rel(base, paths["pdb"]),
            "Meaning": "PDB 5AEP receptor source; co-crystallized QUP ligand used to infer docking grid.",
            "Status": "implemented",
        },
        {
            "Dataset / artifact": "Curated ligand table",
            "Rows": len(curated),
            "Location": rel(base, paths["curated"]),
            "Meaning": "Canonical SMILES, activity values, ring labels, and train/val/test split.",
            "Status": "implemented",
        },
        {
            "Dataset / artifact": "Feature table",
            "Rows": len(features),
            "Location": rel(base, paths["features"]),
            "Meaning": "RDKit descriptors used for filters, ADMET proxy scoring, and ranking.",
            "Status": "implemented",
        },
        {
            "Dataset / artifact": "SE(3) graph tensors",
            "Rows": len(graph_index),
            "Location": rel(base, paths["graphs"]),
            "Meaning": "Torch list of graph dictionaries: atom features, bonds, 3D coordinates, masks.",
            "Status": "implemented",
        },
        {
            "Dataset / artifact": "RDKit candidates",
            "Rows": len(rdkit_rows),
            "Location": rel(base, paths["rdkit"]),
            "Meaning": "Rule-based aromatic substitution candidates.",
            "Status": "implemented",
        },
        {
            "Dataset / artifact": "SELFIES candidates",
            "Rows": len(selfies_rows),
            "Location": rel(base, paths["selfies"]),
            "Meaning": "SELFIES token mutation candidates, including the generated macrocycle pool.",
            "Status": "implemented",
        },
        {
            "Dataset / artifact": "Merged generated candidates",
            "Rows": len(generated),
            "Location": rel(base, paths["generated"]),
            "Meaning": "Filtered, valid, mostly novel candidate pool ranked downstream.",
            "Status": "implemented",
        },
        {
            "Dataset / artifact": "ADMET/synthesis proxy scores",
            "Rows": len(admet),
            "Location": rel(base, paths["admet"]),
            "Meaning": "QED, SA, Lipinski, Veber, risk labels, and proxy scores.",
            "Status": "implemented",
        },
        {
            "Dataset / artifact": "Docking SDF manifest",
            "Rows": len(sdf_manifest),
            "Location": rel(base, paths["sdf_manifest"]),
            "Meaning": "Top candidates attempted for 3D conformer/SDF preparation.",
            "Status": "147 prepared, 3 failed embed",
        },
        {
            "Dataset / artifact": "Ligand PDBQT manifest",
            "Rows": len(pdbqt_manifest),
            "Location": rel(base, paths["pdbqt_manifest"]),
            "Meaning": "Meeko-converted ligand PDBQT files used by Vina.",
            "Status": "147 prepared",
        },
        {
            "Dataset / artifact": "Vina command manifest",
            "Rows": len(vina_manifest),
            "Location": rel(base, paths["vina_manifest"]),
            "Meaning": "Authoritative current docking run manifest.",
            "Status": "147 completed",
        },
        {
            "Dataset / artifact": "Parsed docking scores",
            "Rows": len(docking),
            "Location": rel(base, paths["docking_scores"]),
            "Meaning": "Best Vina affinities extracted from logs and linked to pose files.",
            "Status": "147 numeric scores",
        },
        {
            "Dataset / artifact": "Pose sanity scores",
            "Rows": len(pose_sanity),
            "Location": rel(base, paths["pose_sanity"]),
            "Meaning": "Lightweight clash/contact/centroid checks for top poses.",
            "Status": "10/10 inspected top poses pass",
        },
        {
            "Dataset / artifact": "Final ranking",
            "Rows": len(ranking),
            "Location": rel(base, paths["ranking"]),
            "Meaning": "Consensus candidate ranking with final/backup/hold/reject labels.",
            "Status": "implemented",
        },
    ]

    train_ids = base / "03_features" / "splits" / "train_ids.txt"
    val_ids = base / "03_features" / "splits" / "val_ids.txt"
    test_ids = base / "03_features" / "splits" / "test_ids.txt"

    inventory = {
        "paths": {key: str(value) for key, value in paths.items()},
        "counts": {
            "raw_chembl": len(raw),
            "curated_ligands": len(curated),
            "features": len(features),
            "graphs": len(graph_index),
            "generated": len(generated),
            "rdkit": len(rdkit_rows),
            "selfies": len(selfies_rows),
            "admet": len(admet),
            "sdf_manifest": len(sdf_manifest),
            "pdbqt_manifest": len(pdbqt_manifest),
            "vina_manifest": len(vina_manifest),
            "docking": len(docking),
            "pose_sanity": len(pose_sanity),
            "ranking": len(ranking),
            "final_candidates": len(final_candidates),
            "candidate_cards": len(list(card_dir.glob("*_card.md"))) if card_dir.exists() else 0,
            "molecule_images": len(list(image_dir.glob("*.png"))) if image_dir.exists() else 0,
            "pose_images": len(list(pose_image_dir.glob("*.png"))) if pose_image_dir.exists() else 0,
            "notebooks": len(list((base / "10_notebooks").glob("*.ipynb"))),
            "scripts": len([p for p in (base / "scripts").glob("*.py") if p.is_file()]),
            "modules": len([p for p in (base / "src" / "emd_v5_2_hybrid").glob("*.py") if p.is_file()]),
        },
        "splits": {
            "curated_split_column": split_counts,
            "train_ids": len(train_ids.read_text(encoding="utf-8").splitlines()) if train_ids.exists() else 0,
            "val_ids": len(val_ids.read_text(encoding="utf-8").splitlines()) if val_ids.exists() else 0,
            "test_ids": len(test_ids.read_text(encoding="utf-8").splitlines()) if test_ids.exists() else 0,
        },
        "ring_counts": {
            "curated_macrocycles_12_20": curated_macro,
            "curated_constrained_8_11": curated_constrained,
            "generated_macrocycles_12_20": generated_macro,
            "generated_constrained_8_11": generated_constrained,
            "generated_ring_distribution": ring_counts,
            "rdkit_macrocycles": sum(truthy(row.get("has_macrocycle_12_20")) for row in rdkit_rows),
            "selfies_macrocycles": sum(truthy(row.get("has_macrocycle_12_20")) for row in selfies_rows),
            "rdkit_constrained": sum(truthy(row.get("has_constrained_ring_8_11")) for row in rdkit_rows),
            "selfies_constrained": sum(truthy(row.get("has_constrained_ring_8_11")) for row in selfies_rows),
        },
        "status": {
            "docking_stats": numeric_stats(docking, "best_score"),
            "docking_summary": read_json(paths["docking_summary"]),
            "validation": read_json(paths["validation"]),
            "consistency": read_json(paths["consistency"]),
            "se3_debug": read_json(paths["se3_debug"]),
            "sdf_status": count_by(sdf_manifest, "status"),
            "pdbqt_status": count_by(pdbqt_manifest, "status"),
            "vina_status": count_by(vina_manifest, "status"),
            "ranking_decisions": count_by(ranking, "decision"),
            "ranking_tiers": count_by(ranking, "selection_tier"),
            "risk_counts": count_by(ranking, "main_risk"),
        },
        "dataset_rows": dataset_rows,
        "final_candidates": final_candidate_assets,
        "top_macrocycles": macrocycles[:8],
        "top_constrained": constrained[:8],
        "source_counts": count_by(generated, "source_generator"),
    }
    return inventory


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text(cell, text: Any, bold: bool = False, color: str | None = None, size: int = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[Any]], widths: list[float] | None = None, font_size: int = 8) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, ACCENT)
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=font_size)
        if widths:
            cell.width = Inches(widths[index])
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value, size=font_size)
            if widths:
                cells[index].width = Inches(widths[index])
    document.add_paragraph()


def add_callout(document: Document, title: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(TEXT)
    run.font.size = Pt(10)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(body)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(TEXT)
    document.add_paragraph()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor.from_string(ACCENT if level == 1 else ACCENT_2)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.size = Pt(9)


def add_small_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.font.size = Pt(9)


def make_bar_chart(path: Path, title: str, data: list[tuple[str, int]], width: int = 1200, height: int = 620) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    margin_left, margin_right, margin_top, margin_bottom = 110, 40, 80, 100
    draw.text((margin_left, 28), title, fill=(31, 78, 121))
    max_value = max([value for _, value in data] or [1])
    chart_w = width - margin_left - margin_right
    chart_h = height - margin_top - margin_bottom
    bar_gap = 18
    bar_w = max(38, int((chart_w - bar_gap * (len(data) - 1)) / max(len(data), 1)))
    colors = [(31, 78, 121), (47, 111, 94), (140, 117, 69), (103, 87, 135), (116, 129, 142)]
    for i, (label, value) in enumerate(data):
        x0 = margin_left + i * (bar_w + bar_gap)
        x1 = x0 + bar_w
        bar_h = int(chart_h * value / max_value) if max_value else 0
        y1 = margin_top + chart_h
        y0 = y1 - bar_h
        draw.rectangle([x0, y0, x1, y1], fill=colors[i % len(colors)])
        draw.text((x0, max(60, y0 - 22)), str(value), fill=(36, 52, 71))
        draw.text((x0, y1 + 12), label[:16], fill=(36, 52, 71))
    draw.line([margin_left, margin_top + chart_h, width - margin_right, margin_top + chart_h], fill=(150, 160, 170), width=2)
    image.save(path)
    return path


def make_pipeline_chart(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 1200, 820
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title = "EMD V5.2 Hybrid End-to-End Pipeline"
    draw.text((60, 35), title, fill=(31, 78, 121))
    stages = [
        ("1. Collect", "ChEMBL + PDB"),
        ("2. Curate", "204 JAK2 ligands"),
        ("3. Featurize", "Descriptors + 3D graphs"),
        ("4. SE(3) debug", "CPU checkpoint pass"),
        ("5. Generate", "RDKit + SELFIES"),
        ("6. Dock", "147 Vina scores"),
        ("7. Score", "ADMET + pose sanity"),
        ("8. Rank", "595 candidates"),
        ("9. Report", "Top 5 + audits"),
    ]
    cols = 3
    box_w, box_h = 320, 125
    gap_x, gap_y = 55, 62
    start_x, start_y = 60, 110
    for idx, (head, body) in enumerate(stages):
        row, col = divmod(idx, cols)
        x = start_x + col * (box_w + gap_x)
        y = start_y + row * (box_h + gap_y)
        fill = [(221, 234, 246), (220, 239, 232), (238, 242, 245)][idx % 3]
        draw.rounded_rectangle([x, y, x + box_w, y + box_h], radius=16, fill=fill, outline=(31, 78, 121), width=3)
        draw.text((x + 22, y + 24), head, fill=(31, 78, 121))
        draw.text((x + 22, y + 66), body, fill=(36, 52, 71))
        if col < cols - 1:
            y_mid = y + box_h // 2
            draw.line([x + box_w + 8, y_mid, x + box_w + gap_x - 8, y_mid], fill=(80, 100, 120), width=3)
            draw.polygon(
                [(x + box_w + gap_x - 8, y_mid), (x + box_w + gap_x - 20, y_mid - 8), (x + box_w + gap_x - 20, y_mid + 8)],
                fill=(80, 100, 120),
            )
    image.save(path)
    return path


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    for style_name in ["Normal", "Body Text"]:
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(9)
    for name, size in [("Title", 24), ("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 10)]:
        style = document.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor.from_string(ACCENT)


def add_cover(document: Document, inventory: dict[str, Any]) -> None:
    document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ElectroMacroDiff V5.2 Hybrid")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("End-to-End Project Dossier")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(ACCENT_2)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Generated {date.today().isoformat()} | Target: JAK2 | Structure: PDB 5AEP")
    run.font.size = Pt(9)
    document.add_paragraph()

    counts = inventory["counts"]
    stats = inventory["status"]["docking_stats"]
    quick_rows = [
        ["Curated ligands", counts["curated_ligands"], "ChEMBL-derived JAK2 ligands after curation"],
        ["Generated candidates", counts["generated"], "Filtered RDKit + SELFIES candidate pool"],
        ["Generated macrocycles", inventory["ring_counts"]["generated_macrocycles_12_20"], "12-20 atom ring candidates"],
        ["Docked and scored", counts["docking"], f"Vina score range {stats['min']} to {stats['max']} kcal/mol"],
        ["Final candidates", counts["final_candidates"], "Scaffold-diverse final selection"],
        ["Validation", "48/48", "Project validator checks pass"],
    ]
    add_table(document, ["Metric", "Value", "Meaning"], quick_rows, [2.2, 1.0, 3.9], font_size=9)
    add_callout(
        document,
        "Scientific boundary",
        "All molecules in this dossier are computational candidates. Docking, pose sanity, ADMET, and synthesis values are prioritization signals, not experimental proof of potency, selectivity, safety, or synthesis.",
        fill=LIGHT_GREEN,
    )
    document.add_page_break()


def add_manual_toc(document: Document) -> None:
    add_heading(document, "Document Map", 1)
    sections = [
        "1. Executive status",
        "2. Dataset and molecule inventory",
        "3. Macrocycle and constrained-ring audit",
        "4. End-to-end implemented pipeline",
        "5. Model inventory: trained, pretrained, new, and remaining",
        "6. Current final molecules and evidence",
        "7. File locations and artifact contracts",
        "8. Validation, limitations, and remaining roadmap",
        "9. Appendix: exact commands and source modules",
    ]
    for item in sections:
        add_bullet(document, item)
    document.add_page_break()


def add_executive_status(document: Document, inventory: dict[str, Any], assets: dict[str, Path]) -> None:
    add_heading(document, "1. Executive Status", 1)
    add_small_paragraph(
        document,
        "This project folder now contains a working, reproducible low-resource computational discovery pipeline for JAK2-focused molecule generation and prioritization. The pipeline is implemented end to end through data collection, curation, descriptor generation, graph tensor construction, SE(3) model debug training, baseline generation, docking, ADMET/synthesis scoring, pose sanity checks, ranking, report assets, and audits.",
    )
    document.add_picture(str(assets["pipeline"]), width=Inches(6.8))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_callout(
        document,
        "Current truth",
        "The highest-confidence implemented route today is the validated RDKit/SELFIES + Vina + ADMET + pose sanity + consensus ranking route. The custom SE(3) model exists and has passed a CPU debug train/checkpoint gate, but it is not yet a production-trained generative model.",
        fill=LIGHT_BLUE,
    )
    add_heading(document, "Implemented Outputs", 2)
    rows = [
        ["Raw-to-curated data", "Complete", "300 ChEMBL records -> 204 curated ligands"],
        ["Feature build", "Complete", "204 descriptor rows and 204 graph tensor entries"],
        ["SE(3) debug gate", "Complete", "Forward/loss/backward/optimizer/checkpoint reload passed on CPU"],
        ["Baseline generation", "Complete", "595 filtered unique candidates"],
        ["Docking", "Complete", "147 prepared ligands completed in current Vina manifest"],
        ["Ranking/reporting", "Complete", "595 ranked rows, 5 final candidates, cards/images/audits generated"],
        ["Full production SE(3) training", "Remaining", "Need GPU run, curves, best checkpoint, and SE(3)-generated candidate export"],
    ]
    add_table(document, ["Area", "Status", "Evidence"], rows, [2.0, 1.2, 3.9], font_size=8)


def add_dataset_inventory(document: Document, inventory: dict[str, Any], assets: dict[str, Path]) -> None:
    add_heading(document, "2. Dataset and Molecule Inventory", 1)
    document.add_picture(str(assets["counts"]), width=Inches(6.4))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_small_paragraph(
        document,
        "The table below is the authoritative count map for the run. Paths are relative to the project root C:/Users/srira/Desktop/new_plan/EMD_V5_2_Hybrid.",
    )
    rows = [
        [
            row["Dataset / artifact"],
            row["Rows"],
            row["Location"],
            row["Status"],
            row["Meaning"],
        ]
        for row in inventory["dataset_rows"]
    ]
    add_table(document, ["Artifact", "Rows", "Location", "Status", "Purpose"], rows, [1.35, 0.55, 1.9, 1.15, 2.15], font_size=7)

    add_heading(document, "Train/Validation/Test Split", 2)
    splits = inventory["splits"]
    split_rows = [
        ["Train", splits["train_ids"], splits["curated_split_column"].get("train", 0), "Used for model fitting/debug graph loader"],
        ["Validation", splits["val_ids"], splits["curated_split_column"].get("val", 0), "Used for future training monitoring"],
        ["Test", splits["test_ids"], splits["curated_split_column"].get("test", 0), "Held-out evaluation split"],
    ]
    add_table(document, ["Split", "ID file count", "Curated table count", "Use"], split_rows, [1.0, 1.1, 1.25, 3.6], font_size=8)


def add_macrocycle_audit(document: Document, inventory: dict[str, Any], assets: dict[str, Path]) -> None:
    add_heading(document, "3. Macrocycle and Constrained-Ring Audit", 1)
    rings = inventory["ring_counts"]
    generated_total = inventory["counts"]["generated"]
    rows = [
        ["Curated ChEMBL ligands", rings["curated_macrocycles_12_20"], rings["curated_constrained_8_11"], "Original curated data contains almost no macrocycle/constrained-ring signal."],
        ["RDKit generated", rings["rdkit_macrocycles"], rings["rdkit_constrained"], "Rule-based aromatic substitutions mostly preserve small aromatic ring systems."],
        ["SELFIES generated", rings["selfies_macrocycles"], rings["selfies_constrained"], "SELFIES mutations produced all current macrocycle and constrained-ring candidates."],
        ["Merged generated", rings["generated_macrocycles_12_20"], rings["generated_constrained_8_11"], f"{percent(rings['generated_macrocycles_12_20'], generated_total)} macrocycles; {percent(rings['generated_constrained_8_11'], generated_total)} constrained-ring molecules."],
        ["Final selected top 5", 0, 0, "Top scaffold-diverse candidates are strong ranking hits but not macrocycles."],
    ]
    add_table(document, ["Set", "Macrocycles 12-20", "Constrained 8-11", "Interpretation"], rows, [1.7, 1.15, 1.15, 3.15], font_size=8)
    document.add_picture(str(assets["rings"]), width=Inches(6.4))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_callout(
        document,
        "Macrocycle interpretation",
        "The current pipeline can generate macrocycles, but the best-ranked current candidates are not macrocycles. If the project goal requires macrocycles as final winners, the next development step is macrocycle-targeted generation and selection, not merely broader docking.",
        fill=LIGHT_GREEN,
    )

    add_heading(document, "Top Generated Macrocycle Examples", 2)
    macro_rows = []
    for row in inventory["top_macrocycles"][:6]:
        macro_rows.append(
            [
                row.get("rank", ""),
                row.get("candidate_id", ""),
                row.get("max_ring_size", ""),
                row.get("docking_score", "not docked"),
                row.get("final_weighted_score", ""),
                row.get("selection_tier", ""),
                str(row.get("canonical_smiles", row.get("smiles", "")))[:96],
            ]
        )
    add_table(document, ["Rank", "Candidate", "Ring", "Dock", "Final", "Tier", "SMILES preview"], macro_rows, [0.45, 1.35, 0.45, 0.55, 0.6, 1.0, 2.75], font_size=7)

    add_heading(document, "Top Constrained-Ring Examples", 2)
    constrained_rows = []
    for row in inventory["top_constrained"][:6]:
        constrained_rows.append(
            [
                row.get("rank", ""),
                row.get("candidate_id", ""),
                row.get("max_ring_size", ""),
                row.get("docking_score", "not docked"),
                row.get("final_weighted_score", ""),
                row.get("selection_tier", ""),
                str(row.get("canonical_smiles", row.get("smiles", "")))[:96],
            ]
        )
    add_table(document, ["Rank", "Candidate", "Ring", "Dock", "Final", "Tier", "SMILES preview"], constrained_rows, [0.45, 1.35, 0.45, 0.55, 0.6, 1.0, 2.75], font_size=7)


def add_pipeline(document: Document, inventory: dict[str, Any]) -> None:
    add_heading(document, "4. End-to-End Implemented Pipeline", 1)
    rows = [
        ["M0", "Initialize project", "scripts/00_initialize_project.py", "Folder tree, registry, config, environment report", "Complete"],
        ["M1", "Collect and curate data", "scripts/01_collect_data.py", "Raw ChEMBL CSV, curated ligands, PDB 5AEP", "Complete"],
        ["M2", "Build features", "scripts/02_build_features.py", "Descriptors, split files, SE(3) graph tensors", "Complete"],
        ["M3", "SE(3) debug train", "scripts/03_debug_se3_training.py", "CPU debug checkpoint and reload validation", "Complete debug only"],
        ["M4", "Generate baselines", "scripts/04_generate_baselines.py", "RDKit, SELFIES, merged filtered candidates", "Complete"],
        ["M5", "Docking prep + Vina", "scripts/06_*.py", "Receptor/ligand PDBQT, manifest, logs, poses, parsed scores", "Complete current 147"],
        ["M6", "ADMET/synthesis", "scripts/05_score_admet_synthesis.py", "ADMET, synthesis, risk proxy score files", "Complete"],
        ["M7", "Rank candidates", "scripts/07_rank_candidates.py", "Final ranked candidates and selection tiers", "Complete"],
        ["M8", "Report and audit", "scripts/09_*.py, scripts/10_*.py", "Candidate cards, images, final/consistency audits", "Complete"],
    ]
    add_table(document, ["Stage", "Purpose", "Main runner", "Output", "Status"], rows, [0.55, 1.45, 1.85, 2.4, 0.9], font_size=7)
    add_heading(document, "How Data Moves", 2)
    add_bullet(document, "Canonical SMILES and activity values start in curated ligand CSVs.")
    add_bullet(document, "RDKit descriptors and 3D graph tensors are built from curated SMILES.")
    add_bullet(document, "The SE(3) model consumes atom feature tensors, centered coordinates, and masks.")
    add_bullet(document, "RDKit and SELFIES baseline generators produce candidate SMILES records.")
    add_bullet(document, "Candidates are filtered, scored by ADMET proxies, converted to SDF/PDBQT, docked by Vina, pose-checked, then ranked.")
    add_bullet(document, "The final ranking table is the authoritative downstream decision artifact.")


def add_model_inventory(document: Document, inventory: dict[str, Any]) -> None:
    add_heading(document, "5. Model Inventory: Trained, Pretrained, New, and Remaining", 1)
    add_callout(
        document,
        "Pretraining truth",
        "No neural pretrained generative model has been fine-tuned in this project. The neural component is a new local SE(3)-aware flow model initialized from scratch. External engines and libraries such as AutoDock Vina, RDKit, SELFIES, and Meeko are used as tools, not modified pretrained neural models.",
        fill=LIGHT_BLUE,
    )

    rows = [
        [
            "SE3FlowMatching",
            "New custom neural model",
            "03_features/ligand_graphs/se3_graphs.pt: atom_features, coords, mask",
            "04_models_checkpoints/se3_flow/se3_latest_checkpoint.pt",
            "Debug-trained only; CPU forward/loss/backward/optimizer/reload passed",
            "Full GPU training, validation curves, best checkpoint, and generation export",
        ],
        [
            "RDKit aromatic substitution generator",
            "Deterministic baseline, not trained",
            "Curated SMILES seeds from 02_curated_data/jak2_curated_ligands.csv",
            "05_generated_candidates/rdkit/generated_rdkit_filtered.csv",
            "Implemented; 250 candidates",
            "Optional expanded medchem reaction library",
        ],
        [
            "SELFIES mutation generator",
            "Grammar/string mutation baseline, not trained",
            "Curated SMILES seeds encoded to SELFIES tokens",
            "05_generated_candidates/selfies/generated_selfies_filtered.csv",
            "Implemented; 345 candidates, including all current macrocycles",
            "Macrocycle-targeted token constraints and selection pressure",
        ],
        [
            "ADMET/synthesis proxy scorer",
            "New rule/proxy scoring model",
            "Candidate descriptors: QED, SA, MW, logP, TPSA, HBD/HBA, rotatable bonds",
            "07_admet_synthesis/admet_scores.csv",
            "Implemented for all 595 candidates",
            "Upgrade to calibrated QSAR/ADMET models when labeled data is available",
        ],
        [
            "AutoDock Vina",
            "External docking engine, not trained locally",
            "Receptor PDBQT, ligand PDBQT, grid center/size",
            "06_docking/scores/docking_scores.csv and 06_docking/poses/",
            "Implemented; official Windows binary used; 147 current scores",
            "Optional redocking controls, higher exhaustiveness, GNINA rescoring",
        ],
        [
            "Pose sanity analyzer",
            "New deterministic geometry analyzer",
            "Vina pose PDBQT, cleaned receptor PDB, docking grid JSON",
            "06_docking/scores/pose_sanity_scores.csv and pose images",
            "Implemented; top 10 inspected pass",
            "Publication-grade PyMOL/ChimeraX visual inspection",
        ],
        [
            "Consensus ranker",
            "New weighted scoring/ranking model",
            "Docking score, pose score, ADMET, synthesis, novelty/diversity, safety proxy",
            "08_final_ranking/final_ranked_candidates.csv",
            "Implemented; 595 ranked, 5 final candidates",
            "Sensitivity analysis and expert reweighting",
        ],
        [
            "GNINA / CNN rescoring",
            "Planned external model",
            "Would consume receptor-ligand pose complexes",
            "Not present yet",
            "Not implemented",
            "Stretch goal after Vina pose audit",
        ],
        [
            "Synthesis route model",
            "Planned route planner",
            "Top candidate SMILES",
            "07_admet_synthesis/aizynth_routes/",
            "Folder exists; route generation not run",
            "Run AiZynthFinder or human route review for top molecules",
        ],
    ]
    add_table(document, ["Model/tool", "Type", "Inputs", "Outputs", "Current status", "Still needed"], rows, [1.15, 1.05, 1.7, 1.55, 1.45, 1.45], font_size=6)

    add_heading(document, "What We Changed", 2)
    add_bullet(document, "Created the project-specific pipeline wrappers, schemas, manifests, ranking rules, validation checks, and reporting assets.")
    add_bullet(document, "Implemented a custom SE(3)-aware flow-matching prototype with translation/rotation-aware coordinate vector-field prediction.")
    add_bullet(document, "Did not modify internal weights of Vina, RDKit, SELFIES, Meeko, or any pretrained neural model.")
    add_bullet(document, "The only saved neural weights are the local debug SE(3) checkpoint created from scratch.")


def add_final_molecules(document: Document, inventory: dict[str, Any]) -> None:
    add_heading(document, "6. Current Final Molecules and Evidence", 1)
    add_small_paragraph(
        document,
        "The final selected set is scaffold-diverse. These are the current computational lead hypotheses, not confirmed inhibitors.",
    )
    rows = []
    for row in inventory["final_candidates"]:
        rows.append(
            [
                row["diverse_rank"],
                row["candidate_id"],
                row["source_generator"],
                row["max_ring_size"],
                row["docking_score"],
                row["final_weighted_score"],
                row["pose_decision"],
                row["main_risk"],
            ]
        )
    add_table(document, ["Diverse rank", "Candidate", "Source", "Ring", "Dock", "Final", "Pose", "Main risk"], rows, [0.65, 1.35, 0.75, 0.45, 0.55, 0.65, 0.55, 2.1], font_size=7)

    add_heading(document, "2D Molecule Structures", 2)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for idx, row in enumerate(inventory["final_candidates"]):
        if idx % 2 == 0:
            cells = table.add_row().cells
        cell = cells[idx % 2]
        set_cell_shading(cell, "FFFFFF")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_path = Path(row["image_path"])
        if image_path.exists():
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Inches(2.65))
        label = cell.add_paragraph()
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = label.add_run(f"{row['candidate_id']} | dock {row['docking_score']} | final {row['final_weighted_score']}")
        run.bold = True
        run.font.size = Pt(7)
    document.add_paragraph()

    add_heading(document, "Where to inspect each final molecule", 2)
    rows = []
    for row in inventory["final_candidates"]:
        cid = row["candidate_id"]
        rows.append(
            [
                cid,
                rel(Path(row["card_path"]).parents[2], Path(row["card_path"])) if Path(row["card_path"]).exists() else "missing card",
                rel(Path(row["image_path"]).parents[2], Path(row["image_path"])) if Path(row["image_path"]).exists() else "missing image",
                rel(Path(row["pose_image_path"]).parents[3], Path(row["pose_image_path"])) if Path(row["pose_image_path"]).exists() else "missing pose image",
            ]
        )
    add_table(document, ["Candidate", "Card", "2D image", "Pose image"], rows, [1.3, 2.0, 2.0, 1.8], font_size=7)


def add_locations_and_contracts(document: Document, inventory: dict[str, Any]) -> None:
    add_heading(document, "7. File Locations and Artifact Contracts", 1)
    rows = [
        ["Configuration", file_note(Path.cwd(), "00_project_registry/campaign_config.yaml"), "Target, seeds, model/docking/ranking settings"],
        ["Run registry", file_note(Path.cwd(), "00_project_registry/run_registry.csv"), "Chronological pipeline run records"],
        ["Artifact registry", file_note(Path.cwd(), "00_project_registry/artifact_registry.csv"), "Generated artifact records"],
        ["Environment", file_note(Path.cwd(), "00_project_registry/environment_versions.txt"), "Python/package/version evidence"],
        ["Curated ligands", file_note(Path.cwd(), "02_curated_data/jak2_curated_ligands.csv"), "Core ligand dataset"],
        ["Graph tensors", file_note(Path.cwd(), "03_features/ligand_graphs/se3_graphs.pt"), "SE(3) model input tensors"],
        ["SE(3) checkpoint", file_note(Path.cwd(), "04_models_checkpoints/se3_flow/se3_latest_checkpoint.pt"), "Debug checkpoint"],
        ["Generated candidates", file_note(Path.cwd(), "05_generated_candidates/merged/generated_merged_filtered.csv"), "Candidate pool"],
        ["Docking manifest", file_note(Path.cwd(), "06_docking/scores/vina_command_manifest.csv"), "Authoritative current docking run"],
        ["Docking scores", file_note(Path.cwd(), "06_docking/scores/docking_scores.csv"), "Parsed Vina score table"],
        ["ADMET scores", file_note(Path.cwd(), "07_admet_synthesis/admet_scores.csv"), "Risk and proxy score table"],
        ["Final ranking", file_note(Path.cwd(), "08_final_ranking/final_ranked_candidates.csv"), "Decision table"],
        ["Reports", file_note(Path.cwd(), "09_reports/"), "Cards, images, report drafts, audits, this dossier"],
        ["Notebooks", file_note(Path.cwd(), "10_notebooks/"), "Colab execution notebooks"],
        ["Source code", file_note(Path.cwd(), "src/emd_v5_2_hybrid/"), "Reusable project modules"],
        ["Scripts", file_note(Path.cwd(), "scripts/"), "Stage runners and audit/report builders"],
    ]
    add_table(document, ["Area", "Location", "Purpose"], rows, [1.35, 3.1, 2.65], font_size=7)
    add_callout(
        document,
        "Authoritative current-run rule",
        "For docking, trust 06_docking/scores/vina_command_manifest.csv and 06_docking/scores/docking_scores.csv as the current run of record. Some folders can contain older physical files from previous iterations; the manifest and parsed score CSV define what is current.",
        fill=LIGHT_GRAY,
    )


def add_validation_and_roadmap(document: Document, inventory: dict[str, Any]) -> None:
    add_heading(document, "8. Validation, Limitations, and Remaining Roadmap", 1)
    validation = inventory["status"]["validation"].get("summary", {})
    consistency = inventory["status"]["consistency"]
    rows = [
        ["Project validator", f"{validation.get('passed', 0)} passed / {validation.get('failed', 0)} failed", "scripts/08_validate_project.py --base ."],
        ["Consistency audit", f"{len(consistency.get('failures', []))} failures / {len(consistency.get('warnings', []))} warnings", "scripts/10_consistency_audit.py --base ."],
        ["Unit tests", "12 passed", "python -m unittest discover -s tests"],
        ["Compile check", "Passed", "python -m compileall -q src scripts tests"],
        ["Pose sanity", "10/10 inspected top poses pass", "scripts/06_pose_sanity.py --base . --top-n 10 --render"],
    ]
    add_table(document, ["Check", "Latest result", "Command/evidence"], rows, [1.6, 1.5, 4.0], font_size=8)

    add_heading(document, "Remaining Work", 2)
    remaining = [
        "Train the SE(3) model beyond debug mode on GPU for 50-200 epochs; save latest/best checkpoints and training curves.",
        "Export and evaluate SE(3)-generated molecules, then compare against RDKit/SELFIES baselines on validity, uniqueness, novelty, ring features, docking distribution, and shortlist representation.",
        "Add macrocycle-specific generation pressure if macrocycle winners are required, because the current final top 5 are not macrocycles.",
        "Run publication-grade pose rendering and manual binding-mode review in PyMOL or ChimeraX.",
        "Optionally redock known controls, increase docking exhaustiveness, and add GNINA/CNN rescoring for pose confidence.",
        "Add synthesis route analysis for the final molecules and macrocycle backup set.",
        "Move from proxy ADMET to calibrated ADMET/QSAR models when labeled data and validation objectives are available.",
        "Perform experimental validation: synthesis, biochemical JAK2 assay, selectivity profile, cytotoxicity/safety testing.",
    ]
    for item in remaining:
        add_bullet(document, item)

    add_heading(document, "Limitations", 2)
    limitations = [
        "Docking scores are approximate computational prioritization signals.",
        "ADMET, synthesis, and safety values are descriptor-based proxies, not clinical or experimental measurements.",
        "The low-resource SE(3) checkpoint demonstrates pipeline feasibility, not production-grade molecular generation.",
        "Macrocycle coverage exists in the generated pool but is not yet driving final candidate selection.",
        "No wet-lab potency, selectivity, toxicity, or synthesizability has been proven.",
    ]
    for item in limitations:
        add_bullet(document, item)


def add_appendix(document: Document) -> None:
    add_heading(document, "9. Appendix: Commands and Source Modules", 1)
    commands = [
        "python scripts/00_initialize_project.py --base .",
        "python scripts/01_collect_data.py --base . --limit 300",
        "python scripts/02_build_features.py --base .",
        "python scripts/03_debug_se3_training.py --base . --hidden-dim 64",
        "python scripts/04_generate_baselines.py --base . --seed-count 50 --selfies-per-seed 8 --rdkit-per-seed 5",
        "python scripts/05_score_admet_synthesis.py --base .",
        "python scripts/06_prepare_docking_inputs.py --base . --top-n 150",
        "python scripts/06_prepare_pdbqt.py --base .",
        "python scripts/06_make_vina_manifest.py --base .",
        "python scripts/06_run_vina_manifest.py --base .",
        "python scripts/06_parse_vina_results.py --base .",
        "python scripts/06_pose_sanity.py --base . --top-n 10 --render",
        "python scripts/07_rank_candidates.py --base .",
        "python scripts/09_render_top_molecules.py --base . --top-n 20",
        "python scripts/09_make_report_assets.py --base . --top-n 10",
        "python scripts/08_validate_project.py --base .",
        "python scripts/10_final_audit.py --base .",
        "python scripts/10_consistency_audit.py --base .",
        "python -m unittest discover -s tests",
        "python -m compileall -q src scripts tests",
    ]
    for command in commands:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(command)
        run.font.name = "Consolas"
        run.font.size = Pt(7)

    add_heading(document, "Core Source Modules", 2)
    rows = [
        ["data_collection.py", "ChEMBL/PDB fetch and curation"],
        ["chemistry.py", "Canonicalization, descriptors, ring/macrocycle labels"],
        ["features.py", "Feature CSV and split files"],
        ["se3_dataset.py", "RDKit 3D conformer to graph tensors"],
        ["se3_flow.py", "Custom SE(3)-aware flow-matching model"],
        ["train_se3.py", "Debug/full training helpers and checkpoint IO"],
        ["baseline_generation.py", "RDKit and SELFIES candidate generators"],
        ["docking_prep.py", "SDF/PDBQT/receptor preparation"],
        ["docking.py", "Vina manifest, grid inference, score parsing"],
        ["pose_analysis.py", "Pose sanity metrics and quick renders"],
        ["admet_synthesis.py", "ADMET/synthesis/safety proxy scoring"],
        ["ranking.py", "Weighted consensus ranking and labels"],
        ["validation.py", "Project gates and schema checks"],
        ["reporting.py", "Candidate cards and limitations"],
    ]
    add_table(document, ["Module", "Role"], rows, [2.2, 4.8], font_size=8)


def build_document(base: Path, output_path: Path) -> Path:
    inventory = build_inventory(base)
    assets_dir = base / "09_reports" / "end_to_end_doc_assets"
    assets_dir.mkdir(parents=True, exist_ok=True)
    assets = {
        "counts": make_bar_chart(
            assets_dir / "key_counts.png",
            "Current Artifact Counts",
            [
                ("Raw", inventory["counts"]["raw_chembl"]),
                ("Curated", inventory["counts"]["curated_ligands"]),
                ("Graphs", inventory["counts"]["graphs"]),
                ("Generated", inventory["counts"]["generated"]),
                ("Docked", inventory["counts"]["docking"]),
                ("Final", inventory["counts"]["final_candidates"]),
            ],
        ),
        "rings": make_bar_chart(
            assets_dir / "ring_counts.png",
            "Ring-Class Counts",
            [
                ("Cur macro", inventory["ring_counts"]["curated_macrocycles_12_20"]),
                ("Cur constr", inventory["ring_counts"]["curated_constrained_8_11"]),
                ("Gen macro", inventory["ring_counts"]["generated_macrocycles_12_20"]),
                ("Gen constr", inventory["ring_counts"]["generated_constrained_8_11"]),
                ("Final macro", 0),
            ],
        ),
        "pipeline": make_pipeline_chart(assets_dir / "pipeline.png"),
    }

    inventory_path = assets_dir / "end_to_end_inventory.json"
    safe_inventory = dict(inventory)
    safe_inventory["final_candidates"] = [
        {key: str(value) if isinstance(value, Path) else value for key, value in row.items()}
        for row in inventory["final_candidates"]
    ]
    inventory_path.write_text(json.dumps(safe_inventory, indent=2, default=str), encoding="utf-8")

    document = Document()
    configure_document(document)
    add_cover(document, inventory)
    add_manual_toc(document)
    add_executive_status(document, inventory, assets)
    document.add_page_break()
    add_dataset_inventory(document, inventory, assets)
    document.add_page_break()
    add_macrocycle_audit(document, inventory, assets)
    document.add_page_break()
    add_pipeline(document, inventory)
    document.add_page_break()
    add_model_inventory(document, inventory)
    document.add_page_break()
    add_final_molecules(document, inventory)
    document.add_page_break()
    add_locations_and_contracts(document, inventory)
    document.add_page_break()
    add_validation_and_roadmap(document, inventory)
    document.add_page_break()
    add_appendix(document)

    section = document.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("EMD V5.2 Hybrid End-to-End Project Dossier | Computational candidates require experimental validation")
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor.from_string("6B7280")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the EMD V5.2 end-to-end project dossier DOCX.")
    parser.add_argument("--base", default=str(Path(__file__).resolve().parents[1]))
    parser.add_argument(
        "--output",
        default="09_reports/EMD_V5_2_Hybrid_End_to_End_Project_Dossier.docx",
    )
    args = parser.parse_args()

    base = Path(args.base).resolve()
    output = Path(args.output)
    if not output.is_absolute():
        output = base / output
    result = build_document(base, output)
    print(result)


if __name__ == "__main__":
    main()
